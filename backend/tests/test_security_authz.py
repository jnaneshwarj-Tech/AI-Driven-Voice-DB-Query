"""
Security + authorization hardening tests (Phase 5 of the master brief).

Covers the concrete repairs made during the route audit:
  * Spreadsheet formula injection in CSV/Excel exports is neutralised while the
    visible value stays identical.
  * A client-controlled upload filename can never escape the uploads directory
    (path traversal → arbitrary file read/write).
  * JWT lifecycle: expired / tampered / foreign-signed tokens are rejected, a
    missing secret fails closed, and role enforcement allows the right role and
    rejects the wrong one.
  * The SQL safety validator refuses destructive operations for Staff, so a
    rejected write never reaches MySQL (the record stays unchanged).

Rules honoured: synthetic USNs / in-memory models only, zero production writes,
expected behaviour computed independently of the code under test.
"""
import io
import os
import sys
import csv
import tempfile
import unittest
from datetime import timedelta
from unittest.mock import patch
from types import SimpleNamespace

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from fastapi import HTTPException

import routes_export
import routes_files
import auth
from config import settings
from query_security_validator import validate_sql_query
from auth import require_roles, create_access_token, get_current_user


# ── Helpers ──────────────────────────────────────────────────────────────────
def _table_model(dangerous: str) -> dict:
    """A minimal, valid table-report model with one risky cell (no DB access)."""
    return {
        "title": "Student List",
        "kind": "table",
        "subtitle": "",
        "record_count": 1,
        "institution": routes_export.report_service.INSTITUTION,
        "generated_by": {"username": "tester", "role": "Admin",
                         "timestamp": "2026-01-01 00:00"},
        "notes": [],
        "columns": [{"key": "name", "label": "Name"},
                    {"key": "usn", "label": "USN"}],
        "records": [{"name": dangerous, "usn": "9ZZSEC0001"}],
    }


class FormulaInjectionTests(unittest.TestCase):
    DANGEROUS = [
        "=SUM(1,1)*cmd",
        "+1+1",
        "-2+3",
        "@SUM(1,1)",
        "\t=evil()",
    ]

    def test_helper_prefixes_only_dangerous_leading_chars(self):
        self.assertEqual(routes_export._csv_cell("=x"), "'=x")
        self.assertEqual(routes_export._csv_cell("+x"), "'+x")
        self.assertEqual(routes_export._csv_cell("@x"), "'@x")
        # Ordinary text and numbers must be left byte-for-byte untouched.
        for safe in ("Ramesh", "4HG23CS001", "8.7", "O Positive", "", None, 12, 3.4):
            self.assertEqual(routes_export._csv_cell(safe), safe)

    def test_csv_export_neutralises_formula_cells(self):
        for payload in self.DANGEROUS:
            guarded = payload.lstrip("\t\r")  # \t / \r are triggers but stripped by Excel
            fd, path = tempfile.mkstemp(suffix=".csv")
            os.close(fd)
            try:
                routes_export._build_table_csv(path, _table_model(payload))
                with open(path, newline="", encoding="utf-8-sig") as fh:
                    cells = [c for row in csv.reader(fh) for c in row]
                self.assertIn("'" + guarded, cells,
                              msg=f"no quote guard for {payload!r}: {cells}")
                # No cell may still begin with a formula trigger character.
                self.assertFalse([c for c in cells if c[:1] in ("=", "+", "-", "@")],
                                 msg=f"unguarded formula cell for {payload!r}: {cells}")
            finally:
                os.remove(path)

    def test_excel_export_retypes_formula_cells_to_text(self):
        from openpyxl import load_workbook
        for payload in self.DANGEROUS:
            fd, path = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            try:
                routes_export._build_table_excel(path, _table_model(payload))
                wb = load_workbook(path)
                ws = wb.active
                flat = [c for row in ws.iter_rows() for c in row]
                # No cell may remain classified as a formula.
                self.assertFalse(
                    any(c.data_type == "f" for c in flat if c.value is not None),
                    msg=f"a formula cell survived for {payload!r}")
                # The value is still present as readable text (display unchanged).
                self.assertTrue(
                    any(payload.lstrip("\t\r") in str(c.value) for c in flat),
                    msg=f"payload text lost for {payload!r}")
            finally:
                os.remove(path)


class UploadPathTraversalTests(unittest.TestCase):
    def test_hostile_names_are_contained_inside_uploads(self):
        upload_dir = os.path.realpath(settings.upload_dir_abs)
        hostile = [
            "../../outside.csv", "..\\..\\outside.csv", "/etc/passwd",
            "C:\\Windows\\System32\\evil.dll", "..\\..\\..\\..\\boot.ini",
            os.path.join(upload_dir, "..", "evil.txt"),
            "a/b/c/../../../../escape.csv",
        ]
        for name in hostile:
            resolved = routes_files._upload_path(name)
            self.assertEqual(os.path.dirname(resolved), upload_dir,
                             msg=f"{name!r} escaped to {resolved!r}")
            self.assertTrue(resolved.startswith(upload_dir + os.sep),
                            msg=f"{name!r} not contained: {resolved!r}")

    def test_plain_names_resolve_in_place(self):
        upload_dir = os.path.realpath(settings.upload_dir_abs)
        resolved = routes_files._upload_path("students_2023.csv")
        self.assertEqual(resolved, os.path.join(upload_dir, "students_2023.csv"))

    def test_empty_and_dot_dot_are_rejected(self):
        for bad in ("", "..", "."):
            with self.assertRaises(HTTPException):
                routes_files._upload_path(bad)


class JwtAuthTests(unittest.TestCase):
    SECRET = "unit-test-secret-key-please-ignore"

    def setUp(self):
        self._patch = patch.object(settings, "SECRET_KEY", self.SECRET, create=True)
        self._patch.start()

    def tearDown(self):
        self._patch.stop()

    def test_expired_token_rejected(self):
        tok = create_access_token({"sub": "someone", "role": "Admin"},
                                  expires_delta=timedelta(minutes=-5))
        with self.assertRaises(HTTPException) as ctx:
            get_current_user(tok)
        self.assertEqual(ctx.exception.status_code, 401)

    def test_foreign_signature_rejected(self):
        from jose import jwt
        forged = jwt.encode({"sub": "attacker", "role": "Admin",
                             "exp": 9_999_999_999}, "wrong-secret",
                            algorithm=settings.JWT_ALGORITHM)
        with self.assertRaises(HTTPException) as ctx:
            get_current_user(forged)
        self.assertEqual(ctx.exception.status_code, 401)

    def test_garbage_token_rejected(self):
        with self.assertRaises(HTTPException) as ctx:
            get_current_user("not.a.jwt")
        self.assertEqual(ctx.exception.status_code, 401)

    def test_missing_secret_fails_closed(self):
        with patch.object(settings, "SECRET_KEY", "", create=True):
            with self.assertRaises(HTTPException) as ctx:
                auth._configured_jwt_secret()
            self.assertEqual(ctx.exception.status_code, 503)


class RequireRolesTests(unittest.TestCase):
    def test_allows_configured_role(self):
        dep = require_roles("Staff", "Admin")
        user = {"username": "u", "role": "Staff"}
        self.assertIs(dep(current_user=user), user)

    def test_rejects_wrong_role_with_403(self):
        dep = require_roles("Admin")
        with self.assertRaises(HTTPException) as ctx:
            dep(current_user={"username": "u", "role": "Staff"})
        self.assertEqual(ctx.exception.status_code, 403)

    def test_role_match_is_case_and_whitespace_insensitive(self):
        dep = require_roles("Admin")
        self.assertEqual(dep(current_user={"role": "  ADMIN "})["role"], "  ADMIN ")


class DestructiveSqlGuardTests(unittest.TestCase):
    """A Staff user must never be handed an un-rejected destructive statement,
    so the write is refused before it can reach MySQL — the record is unchanged.
    These assertions use the validator only; no rows are written."""

    def test_staff_delete_is_refused(self):
        result = validate_sql_query(
            {"operation": "delete", "sql": "DELETE FROM students WHERE usn='9ZZSEC0001'"},
            "Staff")
        self.assertFalse(result["is_valid"])

    def test_drop_truncate_delete_all_are_refused(self):
        for sql in ("DROP TABLE students", "TRUNCATE students", "DELETE FROM students"):
            res = validate_sql_query({"operation": "delete", "sql": sql}, "Admin")
            self.assertFalse(res["is_valid"], msg=f"{sql!r} should be refused")


if __name__ == "__main__":
    unittest.main()
