"""
Category-scoped report tests (dashboard "Generate Report").

Rules honored here:
  * Expected values are hand-computed from the synthetic rows we seed (DB
    truth), never from the code under test.
  * Synthetic USNs (9ZZRP…) only; every seeded row is deleted in teardown and
    zero leakage is asserted. Production students are never touched.
  * A category test passes only when the report contains EXACTLY the selected
    records and the category's columns — nothing invented, nothing extra.
"""
import base64
import io
import os
import re
import sys
import unittest
import zlib

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from fastapi import HTTPException

import report_service
import routes_export
from database import db_conn
from routes_export import ReportRequest

ADMIN = {"username": "rp_admin", "role": "Admin"}
STAFF = {"username": "rp_staff", "role": "Staff"}

S1 = "9ZZRP0001"   # discontinued, full fees + marks
S2 = "9ZZRP0002"   # active, unparseable amounts, single paid fee
SGONE = "9ZZRP9999"  # never seeded (missing-record honesty + 404 scope)


def _exec(sql, params=()):
    with db_conn() as conn:
        cur = conn.cursor()
        cur.execute(sql, list(params))
        conn.commit()
        cur.close()


def _seed():
    _cleanup()
    _exec("INSERT INTO students (usn, name, gender, branch, admission_year, student_type, "
          "dob, phone, email, aadhar_no, blood_group, father_name, village, lifecycle_status) "
          "VALUES (%s,'Report One','Male','CSE',2023,'GEN','2005-01-11','9800000001',"
          "'one@college.edu','123456789012','B+','RAVESHA','Holealmanya','Temporarily Discontinued')", (S1,))
    _exec("INSERT INTO students (usn, name, gender, branch, admission_year, student_type, "
          "dob, phone, email, aadhar_no, blood_group, father_name, lifecycle_status) "
          "VALUES (%s,'Report Two','Female','CSE',2023,'BC','2005-06-02','9800000002',"
          "'two@college.edu','223456789012','O','KUMAR','Active')", (S2,))
    _exec("INSERT INTO marks (usn, semester, sgpa, year) VALUES (%s,1,8.0,2024)", (S1,))
    _exec("INSERT INTO marks (usn, semester, sgpa, year) VALUES (%s,2,7.0,2024)", (S1,))
    _exec("INSERT INTO marks (usn, semester, sgpa, year) VALUES (%s,1,6.0,2024)", (S2,))
    _exec("INSERT INTO student_fees (usn, fee_type, year, status, amount_paid, remaining_amount) "
          "VALUES (%s,'VTU',1,'Paid','5000','0')", (S1,))
    _exec("INSERT INTO student_fees (usn, fee_type, year, status, amount_paid, remaining_amount) "
          "VALUES (%s,'PD',1,'Not Paid','0','2500')", (S1,))
    _exec("INSERT INTO student_fees (usn, fee_type, year, status, amount_paid, remaining_amount) "
          "VALUES (%s,'Tuition',1,'Paid','abc','')", (S2,))
    _exec("INSERT INTO student_lifecycle_status (usn, current_status, effective_date, "
          "last_enrolled_semester, scheme, reason, reason_category, document_ref, updated_by) "
          "VALUES (%s,'Temporarily Discontinued','2026-02-01',4,'BE2022',"
          "'Health issues','medical','DOC-9ZZ','rp_admin')", (S1,))


def _cleanup():
    with db_conn() as conn:
        cur = conn.cursor()
        for tbl in ("students", "marks", "student_fees", "student_lifecycle_status",
                    "student_lifecycle_events", "student_backlogs", "student_field_sources"):
            cur.execute(f"DELETE FROM {tbl} WHERE usn IN (%s,%s,%s)", (S1, S2, SGONE))
        conn.commit()
        cur.close()


def _read_response(resp):
    with open(resp.path, "rb") as f:
        raw = f.read()
    try:
        os.remove(resp.path)
    except OSError:
        pass
    return raw


def _inflate_stream(blob: bytes) -> bytes:
    blob = blob.strip()
    if blob[-2:] == b"~>" or blob[:2] == b"<~":
        try:
            blob = base64.a85decode(blob, adobe=True)
        except ValueError:
            pass
    try:
        return zlib.decompress(blob)
    except zlib.error:
        return blob


def _pdf_text(raw: bytes) -> str:
    chunks = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", raw, re.S):
        blob = _inflate_stream(match.group(1))
        for arr in re.finditer(rb"\[(.*?)\]\s*TJ", blob, re.S):
            for tj in re.finditer(rb"\((.*?)\)", arr.group(1), re.S):
                chunks.append(tj.group(1))
        for tj in re.finditer(rb"\((.*?)\)\s*Tj", blob, re.S):
            chunks.append(tj.group(1))
    text = b" ".join(chunks).decode("latin-1", "ignore")
    return re.sub(r"\\(.)", r"\1", text)


def _by_key(model):
    return {c["key"]: c["label"] for c in model["columns"]}


class CategoryModelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _seed()

    @classmethod
    def tearDownClass(cls):
        _cleanup()
        with db_conn() as conn:
            cur = conn.cursor()
            cur.execute("SELECT COUNT(*) FROM students WHERE usn IN (%s,%s,%s)", (S1, S2, SGONE))
            n = cur.fetchone()[0]
            conn.commit()
            cur.close()
        if n != 0:
            raise AssertionError("synthetic report test rows leaked into students")

    def test_student_list_keeps_exactly_what_the_dashboard_shows(self):
        # Spec: the report must contain the displayed records with the same
        # columns, values and order the staff see on screen — no silent
        # substitution of other values, no extra columns.
        rows = [{"usn": S2, "name": "Report Two", "branch": "CSE", "current_sem": 6},
                {"usn": S1, "name": "Report One", "branch": "CSE", "current_sem": 4}]
        m = report_service.build_table_report_model(rows, ADMIN, "student_list")
        self.assertEqual([c["key"] for c in m["columns"]],
                         ["usn", "name", "branch", "current_sem"])
        self.assertEqual(_by_key(m), {"usn": "USN", "name": "Name", "branch": "Branch",
                                      "current_sem": "Semester"})
        recs = m["records"]
        self.assertEqual([r["usn"] for r in recs], [S2, S1])  # dashboard order kept
        self.assertEqual(recs[0]["name"], "Report Two")
        self.assertEqual(recs[1]["current_sem"], 4)

    def test_reports_never_widen_beyond_displayed_records(self):
        m = report_service.build_table_report_model([{"usn": S1}], ADMIN, "student_list")
        self.assertEqual(m["record_count"], 1)
        self.assertEqual(m["records"][0]["usn"], S1)
        with db_conn() as conn:
            cur = conn.cursor()
            cur.execute("SELECT COUNT(*) FROM students")
            total = cur.fetchone()[0]
            cur.close()
        self.assertLess(m["record_count"], total)

    def test_student_list_has_no_fees_or_marks_columns(self):
        m = report_service.build_table_report_model(
            [{"usn": S1, "name": "x"}], ADMIN, "student_list")
        keys = {c["key"] for c in m["columns"]}
        for forbidden in ("sgpa", "cgpa", "fee_status", "amount_paid", "remaining_amount",
                          "backlogs", "aadhar_no"):
            self.assertNotIn(forbidden, keys)

    def test_academic_rows_match_seeded_marks_truth(self):
        # Hand-computed truth: CGPA is the running average of SGPA per semester.
        m = report_service.build_table_report_model(
            [{"usn": S1}, {"usn": S2}], ADMIN, "academic")
        recs = [(r["usn"], r["semester"], r["sgpa"], r["cgpa"]) for r in m["records"]]
        self.assertIn((S1, 1, 8.0, 8.0), recs)
        self.assertIn((S1, 2, 7.0, 7.5), recs)
        self.assertIn((S2, 1, 6.0, 6.0), recs)
        self.assertEqual(len(recs), 3)

    def test_academic_missing_marks_stay_blank_not_invented(self):
        m = report_service.build_table_report_model([{"usn": SGONE}], ADMIN, "academic")
        rec = m["records"][0]
        self.assertEqual((rec["sgpa"], rec["cgpa"]), ("", ""))
        self.assertTrue(any("no recorded marks" in n for n in m["notes"]))

    def test_fees_totals_status_and_honest_blanks(self):
        m = report_service.build_table_report_model(
            [{"usn": S1}, {"usn": S2}, {"usn": SGONE}], ADMIN, "fees")
        by_usn = {r["usn"]: r for r in m["records"]}
        # S1 truth: paid 5000 + 0, pending 0 + 2500 -> mixed -> Partially Paid
        self.assertEqual(by_usn[S1]["total_paid"], 5000)
        self.assertEqual(by_usn[S1]["total_pending"], 2500)
        self.assertEqual(by_usn[S1]["payment_status"], "Partially Paid")
        # S2: unparseable amount -> blank totals (never 0), status from records
        self.assertEqual(by_usn[S2]["total_paid"], "")
        self.assertEqual(by_usn[S2]["total_pending"], "")
        self.assertEqual(by_usn[S2]["payment_status"], "Paid")
        # Unknown student: no fee records -> Not Recorded, never invented zeros
        self.assertEqual(by_usn[SGONE]["payment_status"], "Not Recorded")
        self.assertEqual(by_usn[SGONE]["fee_records"], 0)

    def test_personal_respects_role_masking_and_shows_dob_address(self):
        m = report_service.build_table_report_model([{"usn": S1}], ADMIN, "personal")
        admin_rec = m["records"][0]
        self.assertEqual(admin_rec["aadhar_no"], "123456789012")
        self.assertEqual(admin_rec["dob"], "2005-01-11")
        self.assertEqual(admin_rec["village"], "Holealmanya")
        ms = report_service.build_table_report_model([{"usn": S1}], STAFF, "personal")
        staff_rec = ms["records"][0]
        self.assertNotEqual(staff_rec["aadhar_no"], "123456789012")
        self.assertIn("*", staff_rec["aadhar_no"])
        self.assertIn("*", staff_rec["phone"])
        # Missing data stays blank, not fabricated.
        self.assertEqual(staff_rec["address"], "")

    def test_discontinued_only_includes_bucket_students(self):
        m = report_service.build_table_report_model(
            [{"usn": S1}, {"usn": S2}], ADMIN, "discontinued")
        self.assertEqual(m["record_count"], 1)
        rec = m["records"][0]
        self.assertEqual(rec["usn"], S1)
        self.assertEqual(rec["lifecycle_status"], "Temporarily Discontinued")
        self.assertEqual(str(rec["effective_date"]), "2026-02-01")
        self.assertEqual(rec["reason"], "Health issues")
        self.assertEqual(rec["reason_category"], "medical")
        self.assertTrue(any("excluded" in n for n in m["notes"]))

    def test_category_rejects_unknown(self):
        with self.assertRaises(ValueError):
            report_service.build_table_report_model([{"usn": S1}], ADMIN, "everything")

    def test_student_list_collapses_identical_duplicates_but_keeps_distinct_rows(self):
        # A per-semester join can repeat one student with identical values in
        # every displayed column; those exact duplicates collapse to one line.
        same = [{"usn": S1, "name": "Report One", "current_sem": 4} for _ in range(3)]
        m = report_service.build_table_report_model(same, ADMIN, "student_list")
        self.assertEqual(m["record_count"], 1)
        self.assertTrue(any("collapsed" in n for n in m["notes"]))
        # Rows that differ in any displayed column are ALL kept (no data lost).
        diff = [{"usn": S1, "name": "Report One", "current_sem": 4},
                {"usn": S1, "name": "Report One", "current_sem": 5},
                {"usn": S2, "name": "Report Two", "current_sem": 6}]
        m2 = report_service.build_table_report_model(diff, ADMIN, "student_list")
        self.assertEqual(m2["record_count"], 3)
        self.assertEqual([r["current_sem"] for r in m2["records"]], [4, 5, 6])

    def test_duplicate_join_rows_do_not_repeat_a_student(self):
        # A marks/fees join hands the dashboard the SAME student once per joined
        # row. Student-scoped categories summarize a student, so those duplicate
        # input rows must collapse to a single block — never repeat it. Truth is
        # hand-computed from the seed: S1 has 2 semesters and 1 fee aggregate.
        dup = [{"usn": S1}, {"usn": S1}, {"usn": S1}]
        acad = report_service.build_table_report_model(dup, ADMIN, "academic")
        self.assertEqual(acad["record_count"], 2)  # sem 1 + sem 2, not 6
        fees = report_service.build_table_report_model(dup, ADMIN, "fees")
        self.assertEqual(fees["record_count"], 1)
        self.assertEqual(fees["records"][0]["fee_records"], 2)
        pers = report_service.build_table_report_model(dup, ADMIN, "personal")
        self.assertEqual(pers["record_count"], 1)
        disc = report_service.build_table_report_model(dup, ADMIN, "discontinued")
        self.assertEqual(disc["record_count"], 1)


class ReportEndpointTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _seed()

    @classmethod
    def tearDownClass(cls):
        _cleanup()

    def _req(self, category, fmt="pdf", rows=None):
        return ReportRequest(category=category, format=fmt,
                             rows=rows if rows is not None else
                             [{"usn": S1, "name": "Report One"}, {"usn": S2, "name": "Report Two"}])

    def test_student_list_csv_contains_only_selected_records_and_columns(self):
        raw = _read_response(routes_export.generate_report(
            self._req("student_list", "csv", [{"usn": S1, "name": "Report One"}]),
            ADMIN)).decode("utf-8-sig")
        lines = [ln for ln in raw.splitlines() if ln.strip()]
        self.assertEqual(lines[0], "USN,Name")
        self.assertEqual(len(lines), 2)  # header + the single displayed record
        self.assertIn(S1, raw)
        self.assertNotIn(S2, raw)
        for forbidden in ("FEES", "BACKLOG", "SGPA", "Aadhaar"):
            self.assertNotIn(forbidden, raw.upper())

    def test_excel_has_filter_header_and_widths(self):
        raw = _read_response(routes_export.generate_report(self._req("fees", "excel"), ADMIN))
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw))
        ws = wb.active
        self.assertEqual(ws.title, "Fees Report")
        self.assertIsNotNone(ws.auto_filter.ref)
        header_cells = [c.value for c in ws[ws.auto_filter.ref.split(":")[0][1:]]]
        self.assertEqual(header_cells, ["USN", "Name", "Fee Records", "Total Paid (Rs)",
                                        "Total Pending (Rs)", "Payment Status"])
        self.assertIsNotNone(ws.freeze_panes)
        widths = [ws.column_dimensions[k].width for k in ("A", "B")]
        self.assertTrue(all(w and w > 5 for w in widths))

    def test_pdf_is_a_real_table_report_with_page_numbers_and_no_extra_sections(self):
        raw = _read_response(routes_export.generate_report(self._req("student_list"), ADMIN))
        self.assertEqual(raw[:5], b"%PDF-")
        text = _pdf_text(raw)
        self.assertIn("Student List Report", text)
        self.assertIn(S1, text)
        self.assertIn("Page 1 of", text)
        self.assertNotIn("FEES DETAILS", text)
        self.assertNotIn("BACKLOGS", text)

    def test_print_html_is_category_table(self):
        resp = routes_export.generate_report(self._req("academic", "print"), ADMIN)
        html = resp.body.decode()
        self.assertIn("<table>", html)
        self.assertIn("Cumulative CGPA", html)
        self.assertIn("@page", html)
        self.assertNotIn("FEES DETAILS", html)

    def test_personal_csv_masks_for_staff(self):
        raw = _read_response(routes_export.generate_report(
            self._req("personal", "csv", [{"usn": S1}]), STAFF)).decode("utf-8-sig")
        self.assertIn(S1, raw)
        self.assertNotIn("123456789012", raw)

    def test_individual_requires_exactly_one_student(self):
        with self.assertRaises(HTTPException) as ctx:
            routes_export.generate_report(self._req("individual", "csv"), ADMIN)
        self.assertEqual(ctx.exception.status_code, 400)
        raw = _read_response(routes_export.generate_report(
            self._req("individual", "pdf", [{"usn": S1}]), ADMIN))
        self.assertEqual(raw[:5], b"%PDF-")
        text = _pdf_text(raw)
        # The individual report IS the full authorized profile.
        self.assertIn("FEES DETAILS", text)
        self.assertIn("ACADEMIC PERFORMANCE", text)

    def test_preview_returns_exact_columns_and_truncates(self):
        rows = [{"usn": S1, "name": "Report One", "current_sem": 4}]
        pv = routes_export.report_preview(self._req("student_list", rows=rows), ADMIN)
        self.assertEqual(pv["title"], "Student List Report")
        self.assertEqual([c["key"] for c in pv["columns"]], ["usn", "name", "current_sem"])
        self.assertEqual(pv["record_count"], 1)
        pv2 = routes_export.report_preview(self._req("academic", rows=[{"usn": S1}]), ADMIN)
        self.assertEqual(pv2["record_count"], 2)  # two seeded semesters, from DB truth

    def test_empty_rows_rejected(self):
        with self.assertRaises(HTTPException) as ctx:
            routes_export.generate_report(self._req("student_list", rows=[]), ADMIN)
        self.assertEqual(ctx.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main(verbosity=2)
